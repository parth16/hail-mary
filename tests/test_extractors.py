from __future__ import annotations

import csv
import zipfile
from pathlib import Path

import pytest
from docx import Document

from hailmary.ingest import extractors
from hailmary.ingest.extractors import extract_document
from hailmary.ingest.ocr import LocalOcrError, LocalOcrResult, SubprocessLocalOcrEngine
from hailmary.schemas.documents import ExtractedPage, ExtractionQuality


class FakeOcrEngine:
    def __init__(
        self,
        *,
        image_result: LocalOcrResult | None = None,
        pdf_results: dict[int, LocalOcrResult] | None = None,
        pdf_errors: dict[int, LocalOcrError] | None = None,
    ) -> None:
        self.image_result = image_result or LocalOcrResult(text="")
        self.pdf_results = pdf_results or {}
        self.pdf_errors = pdf_errors or {}
        self.image_calls: list[Path] = []
        self.pdf_calls: list[int] = []

    def image_to_text(
        self,
        path: Path,
        *,
        page_number: int | None = None,
    ) -> LocalOcrResult:
        assert page_number == 1
        self.image_calls.append(path)
        return self.image_result

    def pdf_page_to_text(self, path: Path, *, page_number: int) -> LocalOcrResult:
        del path
        self.pdf_calls.append(page_number)
        error = self.pdf_errors.get(page_number)
        if error is not None:
            raise error
        return self.pdf_results.get(page_number, LocalOcrResult(text=""))


def test_html_extraction_removes_markup_and_scripts(tmp_path: Path) -> None:
    html_path = tmp_path / "deal.html"
    html_path.write_text(
        """
        <html>
          <head><script>window.secret = "ignore";</script></head>
          <body><h1>Acme Deal</h1><p>Customer traction is growing.</p></body>
        </html>
        """,
        encoding="utf-8",
    )

    result = extract_document(html_path)

    assert result.pages
    assert "Acme Deal" in result.combined_text
    assert "Customer traction is growing." in result.combined_text
    assert "script" not in result.combined_text.lower()
    assert "window.secret" not in result.combined_text


def test_pdf_page_access_failure_is_recorded(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pdf_path = tmp_path / "locked.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")

    class BrokenReader:
        @property
        def pages(self) -> object:
            raise RuntimeError("PDF is locked")

    monkeypatch.setattr(extractors, "PdfReader", lambda _: BrokenReader())

    result = extract_document(pdf_path)

    assert result.extraction_quality == ExtractionQuality.LOW
    assert result.notes is not None
    assert "Could not read pages" in result.notes


def test_pdf_page_extraction_failure_is_recorded_on_document(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pdf_path = tmp_path / "partial.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")

    class BrokenPage:
        def extract_text(self) -> str:
            raise RuntimeError("page text is unavailable")

    class GoodPage:
        def extract_text(self) -> str:
            return "Customer traction is strong."

    class PartialReader:
        pages = [BrokenPage(), GoodPage()]

    monkeypatch.setattr(extractors, "PdfReader", lambda _: PartialReader())

    result = extract_document(pdf_path)

    assert result.pages[0].notes is not None
    assert "Could not extract text from page 1" in result.pages[0].notes
    assert result.notes is not None
    assert "Could not extract text from page 1" in result.notes


def test_pdf_quality_is_capped_when_most_pages_need_ocr() -> None:
    pages = [
        ExtractedPage(
            page_number=1,
            raw_text=" ".join(["traction"] * 220),
            clean_text=" ".join(["traction"] * 220),
            word_count=220,
            needs_ocr=False,
        ),
        ExtractedPage(page_number=2, raw_text="", clean_text="", word_count=0, needs_ocr=True),
        ExtractedPage(page_number=3, raw_text="", clean_text="", word_count=0, needs_ocr=True),
        ExtractedPage(page_number=4, raw_text="", clean_text="", word_count=0, needs_ocr=True),
    ]

    assert extractors._quality_from_pages(pages) == ExtractionQuality.MEDIUM


def test_docx_text_traversal_failure_is_recorded(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    docx_path = tmp_path / "broken.docx"
    docx_path.write_bytes(b"placeholder")

    def fail_text_parts(document: object) -> list[str]:
        raise RuntimeError("document part is malformed")

    monkeypatch.setattr(extractors, "Document", lambda _: object())
    monkeypatch.setattr(extractors, "_docx_text_parts", fail_text_parts)

    result = extract_document(docx_path)

    assert result.pages == []
    assert result.extraction_quality == ExtractionQuality.LOW
    assert result.notes is not None
    assert "Could not extract text from the DOCX file" in result.notes


def test_missing_text_file_is_recorded_without_crashing(tmp_path: Path) -> None:
    result = extract_document(tmp_path / "missing.txt")

    assert result.pages == []
    assert result.extraction_quality == ExtractionQuality.LOW
    assert result.notes is not None
    assert "Could not read the text file" in result.notes


def test_raw_only_text_page_is_kept_after_watermark_cleanup(tmp_path: Path) -> None:
    text_path = tmp_path / "watermark.txt"
    text_path.write_text(
        "Not for distribution\nNot for distribution\nNot for distribution\n",
        encoding="utf-8",
    )

    result = extract_document(text_path)

    assert result.pages
    assert result.pages[0].clean_text == ""
    assert result.pages[0].removed_boilerplate_lines == 3
    assert "Not for distribution" in result.combined_raw_text


def test_image_extraction_records_vision_needed_metadata(tmp_path: Path) -> None:
    image_path = tmp_path / "scan.png"
    image_path.write_bytes(b"synthetic image placeholder")

    result = extract_document(image_path)

    assert result.extraction_quality == ExtractionQuality.LOW
    assert result.ocr_recommended
    assert result.vision_recommended
    assert result.page_count == 1
    assert result.pages[0].needs_ocr
    assert result.pages[0].vision_recommended
    assert result.pages[0].source_span_start == 0
    assert result.pages[0].source_span_end == 0
    assert result.combined_text == ""
    assert result.notes is not None
    assert "local OCR" in result.notes


def test_image_ocr_success_creates_page_text(tmp_path: Path) -> None:
    image_path = tmp_path / "scan.png"
    image_path.write_bytes(b"synthetic image placeholder")
    engine = FakeOcrEngine(
        image_result=LocalOcrResult(
            text="Valuation cap $8M. Customer traction is growing.",
            confidence=0.91,
        )
    )

    result = extract_document(image_path, ocr_engine=engine)

    assert result.ocr_applied
    assert result.ocr_confidence == 0.91
    assert not result.ocr_recommended
    assert not result.vision_recommended
    assert result.extraction_quality == ExtractionQuality.MEDIUM
    assert result.combined_text == "Valuation cap $8M. Customer traction is growing."
    assert result.pages[0].page_number == 1
    assert result.pages[0].raw_text == "Valuation cap $8M. Customer traction is growing."
    assert result.pages[0].source_span_start == 0
    assert result.pages[0].source_span_end == len(result.pages[0].raw_text)
    assert result.pages[0].ocr_applied
    assert result.pages[0].ocr_confidence == 0.91
    assert not result.pages[0].needs_ocr
    assert result.pages[0].notes is not None
    assert "Image-based text reading (OCR) was used" in result.pages[0].notes


def test_image_ocr_missing_tesseract_warns_without_crashing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    image_path = tmp_path / "scan.png"
    image_path.write_bytes(b"synthetic image placeholder")
    monkeypatch.setattr("hailmary.ingest.ocr.shutil.which", lambda _: None)

    result = extract_document(image_path, ocr_engine=SubprocessLocalOcrEngine())

    assert result.extraction_quality == ExtractionQuality.LOW
    assert result.ocr_recommended
    assert result.vision_recommended
    assert not result.ocr_applied
    assert result.pages[0].needs_ocr
    assert result.notes is not None
    assert "tesseract" in result.notes
    assert "OCR means reading text from images" in result.notes


def test_image_ocr_low_confidence_is_recorded(tmp_path: Path) -> None:
    image_path = tmp_path / "scan.jpg"
    image_path.write_bytes(b"synthetic image placeholder")
    engine = FakeOcrEngine(
        image_result=LocalOcrResult(
            text="Minimum investment $1,000.",
            confidence=0.24,
        )
    )

    result = extract_document(image_path, ocr_engine=engine)

    assert result.ocr_applied
    assert result.ocr_confidence == 0.24
    assert result.notes is not None
    assert "low confidence" in result.notes
    assert result.pages[0].notes is not None
    assert "low confidence" in result.pages[0].notes


def test_missing_image_file_is_recorded_without_crashing(tmp_path: Path) -> None:
    result = extract_document(tmp_path / "missing.png")

    assert result.pages == []
    assert result.extraction_quality == ExtractionQuality.LOW
    assert result.notes is not None
    assert "Could not read the image file" in result.notes


def test_pdf_empty_page_is_marked_for_ocr_and_vision(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pdf_path = tmp_path / "scan.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")

    class EmptyPage:
        def extract_text(self) -> str:
            return ""

    class ScannedReader:
        pages = [EmptyPage()]

    monkeypatch.setattr(extractors, "PdfReader", lambda _: ScannedReader())

    result = extract_document(pdf_path)

    assert result.ocr_recommended
    assert result.vision_recommended
    assert result.pages[0].needs_ocr
    assert result.pages[0].vision_recommended
    assert result.pages[0].source_span_start == 0
    assert result.pages[0].source_span_end == 0
    assert result.pages[0].notes is not None
    assert "no extracted text" in result.pages[0].notes
    assert result.notes is not None
    assert "local OCR" in result.notes


def test_pdf_empty_page_uses_fake_ocr_text_and_preserves_page_number(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pdf_path = tmp_path / "scan.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")

    class EmptyPage:
        def extract_text(self) -> str:
            return ""

    class ScannedReader:
        pages = [EmptyPage()]

    monkeypatch.setattr(extractors, "PdfReader", lambda _: ScannedReader())
    engine = FakeOcrEngine(
        pdf_results={
            1: LocalOcrResult(
                text="Valuation cap $8M. Minimum investment $1,000.",
                confidence=0.88,
            )
        }
    )

    result = extract_document(pdf_path, ocr_engine=engine)

    assert engine.pdf_calls == [1]
    assert result.ocr_applied
    assert result.ocr_confidence == 0.88
    assert not result.ocr_recommended
    assert result.pages[0].page_number == 1
    assert result.pages[0].raw_text == "Valuation cap $8M. Minimum investment $1,000."
    assert result.pages[0].source_span_start == 0
    assert result.pages[0].source_span_end == len(result.pages[0].raw_text)
    assert result.pages[0].ocr_applied
    assert result.pages[0].ocr_confidence == 0.88
    assert not result.pages[0].needs_ocr


def test_pdf_partial_ocr_failure_keeps_other_pages(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pdf_path = tmp_path / "partial-scan.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")

    class EmptyPage:
        def extract_text(self) -> str:
            return ""

    class ScannedReader:
        pages = [EmptyPage(), EmptyPage()]

    monkeypatch.setattr(extractors, "PdfReader", lambda _: ScannedReader())
    engine = FakeOcrEngine(
        pdf_results={
            2: LocalOcrResult(
                text="Customer revenue is growing.",
                confidence=0.86,
            )
        },
        pdf_errors={
            1: LocalOcrError(
                "Could not run image-based text reading (OCR) on page 1."
            )
        },
    )

    result = extract_document(pdf_path, ocr_engine=engine)

    assert engine.pdf_calls == [1, 2]
    assert result.ocr_applied
    assert result.ocr_recommended
    assert result.vision_recommended
    assert result.pages[0].needs_ocr
    assert not result.pages[0].ocr_applied
    assert result.pages[0].notes is not None
    assert "Could not run image-based text reading (OCR)" in result.pages[0].notes
    assert result.pages[1].ocr_applied
    assert not result.pages[1].needs_ocr
    assert result.combined_text == "Customer revenue is growing."
    assert result.notes is not None
    assert "Could not run image-based text reading (OCR)" in result.notes


def test_pdf_missing_pdftoppm_warns_without_crashing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pdf_path = tmp_path / "scan.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")

    class EmptyPage:
        def extract_text(self) -> str:
            return ""

    class ScannedReader:
        pages = [EmptyPage()]

    def fake_which(command: str) -> str | None:
        if command == "tesseract":
            return "/usr/local/bin/tesseract"
        return None

    monkeypatch.setattr(extractors, "PdfReader", lambda _: ScannedReader())
    monkeypatch.setattr("hailmary.ingest.ocr.shutil.which", fake_which)

    result = extract_document(pdf_path, ocr_engine=SubprocessLocalOcrEngine())

    assert result.ocr_recommended
    assert result.vision_recommended
    assert not result.ocr_applied
    assert result.pages[0].needs_ocr
    assert result.notes is not None
    assert "pdftoppm" in result.notes
    assert "OCR means reading text from images" in result.notes


def test_pdf_short_text_page_does_not_recommend_ocr(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pdf_path = tmp_path / "deck.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")

    class CoverPage:
        def extract_text(self) -> str:
            return "Acme Investor Deck"

    class Reader:
        pages = [CoverPage()]

    monkeypatch.setattr(extractors, "PdfReader", lambda _: Reader())

    result = extract_document(pdf_path)

    assert not result.ocr_recommended
    assert not result.vision_recommended
    assert not result.pages[0].needs_ocr


def test_pdf_repeated_short_text_pages_recommend_ocr(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pdf_path = tmp_path / "deck.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")

    class TextPage:
        def __init__(self, text: str) -> None:
            self.text = text

        def extract_text(self) -> str:
            return self.text

    class Reader:
        pages = [
            TextPage("Customer logo slide with visual proof"),
            TextPage("Product demo slide with embedded screenshot"),
            TextPage("Market map slide with competitor logos"),
        ]

    monkeypatch.setattr(extractors, "PdfReader", lambda _: Reader())

    result = extract_document(pdf_path)

    assert result.ocr_recommended
    assert result.vision_recommended
    assert [page.needs_ocr for page in result.pages] == [True, True, True]
    assert result.notes is not None
    assert "local OCR" in result.notes


def test_pdf_image_backed_low_text_page_recommends_ocr(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pdf_path = tmp_path / "scan.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")

    class ImagePage:
        images = [object()]

        def extract_text(self) -> str:
            return "1"

    class TextPage:
        def extract_text(self) -> str:
            return "Readable traction text with customer growth and revenue context."

    class Reader:
        pages = [ImagePage(), TextPage()]

    monkeypatch.setattr(extractors, "PdfReader", lambda _: Reader())

    result = extract_document(pdf_path)

    assert result.ocr_recommended
    assert result.vision_recommended
    assert result.pages[0].needs_ocr
    assert result.pages[0].vision_recommended
    assert result.pages[0].notes is not None
    assert "image content" in result.pages[0].notes
    assert result.notes is not None
    assert "local OCR" in result.notes


def test_pdf_low_text_page_recommends_ocr(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pdf_path = tmp_path / "deck.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")

    class PageNumberOnly:
        def extract_text(self) -> str:
            return "1"

    class Reader:
        pages = [PageNumberOnly()]

    monkeypatch.setattr(extractors, "PdfReader", lambda _: Reader())

    result = extract_document(pdf_path)

    assert result.ocr_recommended
    assert result.vision_recommended
    assert result.pages[0].needs_ocr


def test_pdf_low_text_divider_with_readable_page_does_not_warn_document(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pdf_path = tmp_path / "deck.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")

    class TextPage:
        def __init__(self, text: str) -> None:
            self.text = text

        def extract_text(self) -> str:
            return self.text

    class Reader:
        pages = [
            TextPage("1"),
            TextPage("Readable traction text with customer growth and revenue context."),
        ]

    monkeypatch.setattr(extractors, "PdfReader", lambda _: Reader())

    result = extract_document(pdf_path)

    assert result.pages[0].needs_ocr
    assert not result.ocr_recommended
    assert not result.vision_recommended
    assert result.notes is None


def test_pdf_boilerplate_only_divider_with_readable_page_does_not_warn_document(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pdf_path = tmp_path / "deck.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")

    class TextPage:
        def __init__(self, text: str) -> None:
            self.text = text

        def extract_text(self) -> str:
            return self.text

    class Reader:
        pages = [
            TextPage(
                "Not for distribution\n"
                "Not for distribution\n"
                "Not for distribution\n"
            ),
            TextPage("Readable traction text with customer growth and revenue context."),
        ]

    monkeypatch.setattr(extractors, "PdfReader", lambda _: Reader())

    result = extract_document(pdf_path)

    assert result.pages[0].clean_text == ""
    assert result.pages[0].notes is None
    assert not result.ocr_recommended
    assert not result.vision_recommended
    assert result.notes is None


def test_pdf_empty_page_with_readable_page_warns_document(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pdf_path = tmp_path / "deck.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")

    class TextPage:
        def __init__(self, text: str) -> None:
            self.text = text

        def extract_text(self) -> str:
            return self.text

    class Reader:
        pages = [
            TextPage(""),
            TextPage("Readable traction text with customer growth and revenue context."),
        ]

    monkeypatch.setattr(extractors, "PdfReader", lambda _: Reader())

    result = extract_document(pdf_path)

    assert result.ocr_recommended
    assert result.vision_recommended
    assert result.pages[0].notes is not None
    assert "no extracted text" in result.pages[0].notes
    assert result.notes is not None
    assert "local OCR" in result.notes


def test_pdf_mostly_low_text_pages_recommend_ocr(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pdf_path = tmp_path / "deck.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")

    class TextPage:
        def __init__(self, text: str) -> None:
            self.text = text

        def extract_text(self) -> str:
            return self.text

    class Reader:
        pages = [
            TextPage("1"),
            TextPage("2"),
            TextPage("Readable traction text with customer growth and revenue context."),
        ]

    monkeypatch.setattr(extractors, "PdfReader", lambda _: Reader())

    result = extract_document(pdf_path)

    assert result.ocr_recommended
    assert result.vision_recommended


def test_pdf_text_page_records_source_span(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pdf_path = tmp_path / "deck.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")

    class TextPage:
        def __init__(self, text: str) -> None:
            self.text = text

        def extract_text(self) -> str:
            return self.text

    class Reader:
        pages = [
            TextPage("First page has enough traction words to count as readable text."),
            TextPage("Second page also has enough readable words for extraction quality."),
        ]

    monkeypatch.setattr(extractors, "PdfReader", lambda _: Reader())

    result = extract_document(pdf_path)

    assert not result.vision_recommended
    assert result.pages[0].source_span_start == 0
    assert result.pages[0].source_span_end == len(Reader.pages[0].text)
    assert result.pages[1].source_span_start == len(Reader.pages[0].text) + 2


def test_pdf_source_spans_ignore_empty_pages_before_text(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pdf_path = tmp_path / "deck.pdf"
    pdf_path.write_bytes(b"%PDF-1.4")

    class TextPage:
        def __init__(self, text: str) -> None:
            self.text = text

        def extract_text(self) -> str:
            return self.text

    class Reader:
        pages = [
            TextPage(""),
            TextPage("Readable traction text follows the empty cover page."),
        ]

    monkeypatch.setattr(extractors, "PdfReader", lambda _: Reader())

    result = extract_document(pdf_path)

    assert result.combined_raw_text == Reader.pages[1].text
    assert result.pages[1].source_span_start == 0
    assert result.pages[1].source_span_end == len(Reader.pages[1].text)


def test_docx_table_extraction_records_structured_rows(tmp_path: Path) -> None:
    docx_path = tmp_path / "memo.docx"
    document = Document()
    document.add_paragraph("Acme diligence memo.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "ARR"
    table.cell(1, 1).text = "$1M"
    document.save(str(docx_path))

    result = extract_document(docx_path)

    assert result.table_count == 1
    assert result.tables[0].rows == [["Metric", "Value"], ["ARR", "$1M"]]
    assert result.tables[0].row_count == 2
    assert result.tables[0].column_count == 2
    assert "ARR | $1M" in result.combined_text


def test_html_table_extraction_records_structured_rows(tmp_path: Path) -> None:
    html_path = tmp_path / "deal.html"
    html_path.write_text(
        """
        <html><body>
          <table>
            <tr><th>Customer</th><th>Status</th></tr>
            <tr><td>Acme Bank</td><td>Pilot</td></tr>
          </table>
        </body></html>
        """,
        encoding="utf-8",
    )

    result = extract_document(html_path)

    assert result.table_count == 1
    assert result.tables[0].rows == [["Customer", "Status"], ["Acme Bank", "Pilot"]]
    assert "Customer | Status" in result.tables[0].clean_text


def test_html_nested_tables_are_captured_separately(tmp_path: Path) -> None:
    html_path = tmp_path / "deal.html"
    html_path.write_text(
        """
        <html><body>
          <table>
            <tbody>
              <tr><th>Metric</th><th>Detail</th></tr>
              <tr>
                <td>ARR</td>
                <td>
                  <table>
                    <tr><th>Nested</th><th>Value</th></tr>
                    <tr><td>Expansion</td><td>Strong</td></tr>
                  </table>
                </td>
              </tr>
            </tbody>
          </table>
        </body></html>
        """,
        encoding="utf-8",
    )

    result = extract_document(html_path)

    assert result.table_count == 2
    assert result.tables[0].rows == [["Metric", "Detail"], ["ARR"]]
    assert result.tables[1].rows == [["Nested", "Value"], ["Expansion", "Strong"]]


def test_html_table_extraction_expands_row_and_column_spans(tmp_path: Path) -> None:
    html_path = tmp_path / "deal.html"
    html_path.write_text(
        """
        <html><body>
          <table>
            <tr><th rowspan="2">Metric</th><th colspan="2">Revenue</th></tr>
            <tr><th>2025</th><th>2026</th></tr>
            <tr><td>ARR</td><td>$1M</td><td>$2M</td></tr>
          </table>
        </body></html>
        """,
        encoding="utf-8",
    )

    result = extract_document(html_path)

    assert result.table_count == 1
    assert result.tables[0].rows == [
        ["Metric", "Revenue", ""],
        ["Metric", "2025", "2026"],
        ["ARR", "$1M", "$2M"],
    ]
    assert result.tables[0].column_count == 3
    assert "Metric | 2025 | 2026" in result.tables[0].clean_text


def test_csv_extraction_records_table_text(tmp_path: Path) -> None:
    csv_path = tmp_path / "model.csv"
    csv_path.write_text("year,revenue\n2026,100\n", encoding="utf-8")

    result = extract_document(csv_path)

    assert result.pages
    assert "year | revenue" in result.combined_text
    assert "2026 | 100" in result.combined_text
    assert result.table_count == 1
    assert result.tables[0].rows == [["year", "revenue"], ["2026", "100"]]


def test_csv_extraction_preserves_interior_blank_cells(tmp_path: Path) -> None:
    csv_path = tmp_path / "model.csv"
    csv_path.write_text("metric,,value,\nARR,,100,\n", encoding="utf-8")

    result = extract_document(csv_path)

    assert "metric |  | value" in result.tables[0].clean_text
    assert "ARR |  | 100" in result.tables[0].clean_text
    assert result.tables[0].column_count == 3


def test_csv_extraction_filters_blank_rows_from_table_metadata(tmp_path: Path) -> None:
    csv_path = tmp_path / "model.csv"
    csv_path.write_text("metric,value\n,\nARR,100\n", encoding="utf-8")

    result = extract_document(csv_path)

    assert result.tables[0].rows == [["metric", "value"], ["ARR", "100"]]
    assert result.tables[0].row_count == 2


def test_csv_extraction_skips_empty_tables(tmp_path: Path) -> None:
    csv_path = tmp_path / "empty.csv"
    csv_path.write_text(",\n,\n", encoding="utf-8")

    result = extract_document(csv_path)

    assert result.pages == []
    assert result.tables == []
    assert result.table_count == 0


def test_csv_parser_error_is_recorded_without_crashing(tmp_path: Path) -> None:
    csv_path = tmp_path / "model.csv"
    csv_path.write_text(f"notes\n{'A' * 32}\n", encoding="utf-8")
    original_limit = csv.field_size_limit()
    csv.field_size_limit(10)
    try:
        result = extract_document(csv_path)
    finally:
        csv.field_size_limit(original_limit)

    assert result.pages == []
    assert result.extraction_quality == ExtractionQuality.LOW
    assert result.notes is not None
    assert "Could not parse the CSV file" in result.notes


def test_xlsx_extraction_records_table_text(tmp_path: Path) -> None:
    xlsx_path = tmp_path / "model.xlsx"
    with zipfile.ZipFile(xlsx_path, "w") as workbook:
        workbook.writestr(
            "xl/sharedStrings.xml",
            """
            <sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
              <si><t>Revenue</t></si>
            </sst>
            """,
        )
        workbook.writestr(
            "xl/worksheets/sheet1.xml",
            """
            <worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
              <sheetData>
                <row><c t="s"><v>0</v></c><c><v>100</v></c></row>
              </sheetData>
            </worksheet>
            """,
        )

    result = extract_document(xlsx_path)

    assert result.pages
    assert "Revenue | 100" in result.combined_text
    assert result.table_count == 1
    assert result.tables[0].rows == [["Revenue", "100"]]


def test_xlsx_extraction_preserves_sparse_blank_cells(tmp_path: Path) -> None:
    xlsx_path = tmp_path / "model.xlsx"
    with zipfile.ZipFile(xlsx_path, "w") as workbook:
        workbook.writestr(
            "xl/sharedStrings.xml",
            """
            <sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
              <si><t>Metric</t></si>
              <si><t>Value</t></si>
            </sst>
            """,
        )
        workbook.writestr(
            "xl/worksheets/sheet1.xml",
            """
            <worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
              <sheetData>
                <row r="1"><c r="A1" t="s"><v>0</v></c><c r="C1" t="s"><v>1</v></c></row>
              </sheetData>
            </worksheet>
            """,
        )

    result = extract_document(xlsx_path)

    assert result.tables[0].rows == [["Metric", "", "Value"]]
    assert result.tables[0].column_count == 3


def test_xlsx_impossible_cell_reference_is_skipped(tmp_path: Path) -> None:
    xlsx_path = tmp_path / "model.xlsx"
    with zipfile.ZipFile(xlsx_path, "w") as workbook:
        workbook.writestr(
            "xl/sharedStrings.xml",
            """
            <sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
              <si><t>Revenue</t></si>
              <si><t>Impossible</t></si>
            </sst>
            """,
        )
        workbook.writestr(
            "xl/worksheets/sheet1.xml",
            """
            <worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
              <sheetData>
                <row r="1">
                  <c r="A1" t="s"><v>0</v></c>
                  <c r="ZZZZZZ1" t="s"><v>1</v></c>
                </row>
              </sheetData>
            </worksheet>
            """,
        )

    result = extract_document(xlsx_path)

    assert result.tables[0].rows == [["Revenue"]]
    assert result.tables[0].column_count == 1
    assert "Impossible" not in result.combined_text


def test_xlsx_valid_but_huge_sparse_gap_is_compacted(tmp_path: Path) -> None:
    xlsx_path = tmp_path / "model.xlsx"
    with zipfile.ZipFile(xlsx_path, "w") as workbook:
        workbook.writestr(
            "xl/sharedStrings.xml",
            """
            <sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
              <si><t>Revenue</t></si>
              <si><t>Far away</t></si>
            </sst>
            """,
        )
        workbook.writestr(
            "xl/worksheets/sheet1.xml",
            """
            <worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
              <sheetData>
                <row r="1">
                  <c r="A1" t="s"><v>0</v></c>
                  <c r="XFD1" t="s"><v>1</v></c>
                </row>
              </sheetData>
            </worksheet>
            """,
        )

    result = extract_document(xlsx_path)

    assert result.tables[0].rows == [["Revenue", "[16382 blank columns]", "Far away"]]
    assert result.tables[0].column_count == 3
    assert "Far away" in result.combined_text
    assert result.combined_text.count(" | ") == 2


def test_xlsx_empty_far_right_cell_does_not_create_gap_marker(tmp_path: Path) -> None:
    xlsx_path = tmp_path / "model.xlsx"
    with zipfile.ZipFile(xlsx_path, "w") as workbook:
        workbook.writestr(
            "xl/sharedStrings.xml",
            """
            <sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
              <si><t>Revenue</t></si>
            </sst>
            """,
        )
        workbook.writestr(
            "xl/worksheets/sheet1.xml",
            """
            <worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
              <sheetData>
                <row r="1">
                  <c r="A1" t="s"><v>0</v></c>
                  <c r="XFD1"/>
                </row>
              </sheetData>
            </worksheet>
            """,
        )

    result = extract_document(xlsx_path)

    assert result.tables[0].rows == [["Revenue"]]
    assert result.tables[0].column_count == 1
    assert "blank columns" not in result.combined_text


def test_xlsx_compacted_gap_tracks_source_position(tmp_path: Path) -> None:
    xlsx_path = tmp_path / "model.xlsx"
    with zipfile.ZipFile(xlsx_path, "w") as workbook:
        workbook.writestr(
            "xl/sharedStrings.xml",
            """
            <sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
              <si><t>Revenue</t></si>
              <si><t>Far away</t></si>
              <si><t>Adjacent</t></si>
            </sst>
            """,
        )
        workbook.writestr(
            "xl/worksheets/sheet1.xml",
            """
            <worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
              <sheetData>
                <row r="1">
                  <c r="A1" t="s"><v>0</v></c>
                  <c r="DG1" t="s"><v>1</v></c>
                  <c r="DH1" t="s"><v>2</v></c>
                </row>
              </sheetData>
            </worksheet>
            """,
        )

    result = extract_document(xlsx_path)

    assert result.tables[0].rows == [
        ["Revenue", "[109 blank columns]", "Far away", "Adjacent"]
    ]
    assert result.tables[0].column_count == 4
    assert result.combined_text.count("blank columns") == 1


def test_xlsx_negative_shared_string_index_is_left_raw(tmp_path: Path) -> None:
    xlsx_path = tmp_path / "model.xlsx"
    with zipfile.ZipFile(xlsx_path, "w") as workbook:
        workbook.writestr(
            "xl/sharedStrings.xml",
            """
            <sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
              <si><t>Do not use this string</t></si>
            </sst>
            """,
        )
        workbook.writestr(
            "xl/worksheets/sheet1.xml",
            """
            <worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
              <sheetData>
                <row><c t="s"><v>-1</v></c></row>
              </sheetData>
            </worksheet>
            """,
        )

    result = extract_document(xlsx_path)

    assert result.pages
    assert "-1" in result.combined_text
    assert "Do not use this string" not in result.combined_text


def test_xlsx_worksheet_read_failure_is_recorded(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    xlsx_path = tmp_path / "encrypted.xlsx"
    xlsx_path.write_bytes(b"placeholder")

    class BrokenWorkbook:
        def __enter__(self) -> BrokenWorkbook:
            return self

        def __exit__(self, *args: object) -> None:
            return None

        def namelist(self) -> list[str]:
            return ["xl/worksheets/sheet1.xml"]

        def read(self, name: str) -> bytes:
            raise RuntimeError(f"{name} is encrypted")

    monkeypatch.setattr(zipfile, "ZipFile", lambda _: BrokenWorkbook())

    result = extract_document(xlsx_path)

    assert result.pages == []
    assert result.extraction_quality == ExtractionQuality.LOW
    assert result.notes is not None
    assert "Could not read the XLSX file" in result.notes
