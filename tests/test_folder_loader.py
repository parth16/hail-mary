from __future__ import annotations

import hashlib
import json
import os
import stat
import zipfile
from collections.abc import Callable, Iterator
from pathlib import Path

import pytest
from docx import Document

from hailmary.config import AppConfig
from hailmary.ingest import folder_loader
from hailmary.ingest.folder_loader import IngestionError, ingest_folder
from hailmary.schemas.documents import DocumentType, ExtractionQuality, FileType


def _deal_id(deal_name: str) -> str:
    digest = hashlib.sha256(deal_name.encode("utf-8")).hexdigest()[:8]
    slug = "".join(character.lower() if character.isalnum() else "-" for character in deal_name)
    normalized_slug = "-".join(part for part in slug.split("-") if part) or "unknown"
    return f"{normalized_slug}-{digest}"


def _write_docx(path: Path, lines: list[str]) -> None:
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    document.save(str(path))


def test_ingest_folder_groups_documents_and_writes_outputs(tmp_path: Path) -> None:
    root = tmp_path / "pitch-decks"
    company = root / "ExampleCo"
    closing = company / "Closing documents"
    closing.mkdir(parents=True)
    (company / ".DS_Store").write_text("noise", encoding="utf-8")
    (company / "ExampleCo _ AngelList.pdf").write_text("not a real pdf", encoding="utf-8")
    (company / "notes.txt").write_text("Memo about traction and team.", encoding="utf-8")
    _write_docx(
        closing / "ExampleCo - PPM.docx",
        ["PRIVATE PLACEMENT MEMORANDUM", "ExampleCo confidential investment material."],
    )

    config = AppConfig(data_dir=tmp_path / "data")

    summary = ingest_folder(root, config=config)

    assert summary.document_count == 3
    assert len(summary.deals) == 1
    assert summary.deals[0].company_name == "ExampleCo"
    assert ".DS_Store" in " ".join(summary.skipped_files)
    assert summary.summary_path.exists()

    document_types = {doc.source.document_type for doc in summary.deals[0].documents}
    assert DocumentType.PLATFORM_DEAL_PAGE in document_types
    assert DocumentType.LEGAL_DOCUMENT in document_types

    output_paths = [doc.output_path for doc in summary.deals[0].documents]
    assert all(path.exists() for path in output_paths)
    assert summary.deals[0].evidence_store_path is not None
    assert summary.deals[0].evidence_store_path.exists()
    assert summary.deals[0].evidence_count >= 1

    saved_summary = json.loads(summary.summary_path.read_text(encoding="utf-8"))
    assert saved_summary["deals"][0]["id"] == _deal_id("ExampleCo")
    assert saved_summary["deals"][0]["documents"][0]["source"]["company_name"] == "ExampleCo"
    assert saved_summary["deals"][0]["evidence_store_path"]


def test_invalid_pdf_is_recorded_without_crashing(tmp_path: Path) -> None:
    root = tmp_path / "pitch-decks"
    company = root / "BadPdfCo"
    company.mkdir(parents=True)
    (company / "BadPdfCo _ AngelList.pdf").write_text("not a real pdf", encoding="utf-8")

    summary = ingest_folder(root, config=AppConfig(data_dir=tmp_path / "data"))

    document = summary.deals[0].documents[0].source
    assert document.file_type == FileType.PDF
    assert document.document_type == DocumentType.PLATFORM_DEAL_PAGE
    assert document.extraction_quality == ExtractionQuality.LOW
    assert document.notes is not None


def test_confidentiality_marker_survives_watermark_cleanup(tmp_path: Path) -> None:
    root = tmp_path / "pitch-decks"
    company = root / "WatermarkCo"
    company.mkdir(parents=True)
    (company / "memo.txt").write_text(
        "Not for distribution\nNot for distribution\nNot for distribution\n",
        encoding="utf-8",
    )

    summary = ingest_folder(root, config=AppConfig(data_dir=tmp_path / "data"))

    document = summary.deals[0].documents[0]
    assert document.source.confidentiality_detected
    assert document.pages
    assert document.pages[0].clean_text == ""


def test_duplicate_files_get_distinct_output_paths(tmp_path: Path) -> None:
    root = tmp_path / "pitch-decks"
    company = root / "DuplicateCo"
    company.mkdir(parents=True)
    (company / "one.txt").write_text("Same diligence note.", encoding="utf-8")
    (company / "two.txt").write_text("Same diligence note.", encoding="utf-8")

    summary = ingest_folder(root, config=AppConfig(data_dir=tmp_path / "data"))
    output_paths = [doc.output_path for doc in summary.deals[0].documents]

    assert len(output_paths) == 2
    assert len(set(output_paths)) == 2
    assert all(path.exists() for path in output_paths)


def test_direct_company_folder_scan_uses_root_folder_name(tmp_path: Path) -> None:
    root = tmp_path / "Acme"
    nested = root / "Closing Documents"
    root.mkdir()
    nested.mkdir()
    (root / "deck.pdf").write_text("not a real pdf", encoding="utf-8")
    (root / "memo.txt").write_text("Memo about Acme.", encoding="utf-8")
    (nested / "safe.txt").write_text("Simple Agreement for Future Equity.", encoding="utf-8")

    summary = ingest_folder(root, config=AppConfig(data_dir=tmp_path / "data"))

    assert summary.document_count == 3
    assert len(summary.deals) == 1
    assert summary.deals[0].company_name == "Acme"
    assert {document.source.company_name for document in summary.deals[0].documents} == {"Acme"}


def test_top_level_deal_named_data_is_ingested(tmp_path: Path) -> None:
    root = tmp_path / "pitch-decks"
    company = root / "Data"
    company.mkdir(parents=True)
    (company / "memo.txt").write_text("Memo about the Data company.", encoding="utf-8")

    summary = ingest_folder(root, config=AppConfig(data_dir=tmp_path / "generated-data"))

    assert summary.document_count == 1
    assert len(summary.deals) == 1
    assert summary.deals[0].company_name == "Data"


def test_real_nested_data_and_reports_folders_are_ingested(tmp_path: Path) -> None:
    root = tmp_path / "pitch-decks"
    data_folder = root / "Acme" / "Data"
    reports_folder = root / "Acme" / "Reports"
    data_folder.mkdir(parents=True)
    reports_folder.mkdir(parents=True)
    (data_folder / "deck.pdf").write_text("not a real pdf", encoding="utf-8")
    (reports_folder / "memo.txt").write_text("Memo about Acme.", encoding="utf-8")

    summary = ingest_folder(root, config=AppConfig(data_dir=tmp_path / "generated-data"))

    assert summary.document_count == 2


def test_configured_output_folder_inside_scan_root_is_skipped(tmp_path: Path) -> None:
    root = tmp_path / "pitch-decks"
    company = root / "Acme"
    output_folder = root / "generated-data"
    company.mkdir(parents=True)
    output_folder.mkdir(parents=True)
    (company / "memo.txt").write_text("Memo about Acme.", encoding="utf-8")
    (output_folder / "old-output.txt").write_text("This is generated output.", encoding="utf-8")

    summary = ingest_folder(root, config=AppConfig(data_dir=output_folder))

    assert summary.document_count == 1
    assert "generated-data/old-output.txt" in summary.skipped_files


def test_private_raw_folder_can_be_ingested(tmp_path: Path) -> None:
    data_dir = tmp_path / "data"
    raw_company = data_dir / "raw" / "RawCo"
    raw_company.mkdir(parents=True)
    (raw_company / "memo.txt").write_text("Memo about RawCo.", encoding="utf-8")
    processed_folder = data_dir / "processed"
    processed_folder.mkdir()
    (processed_folder / "old-output.txt").write_text("Generated output.", encoding="utf-8")

    summary = ingest_folder(data_dir / "raw", config=AppConfig(data_dir=data_dir))

    assert summary.document_count == 1
    assert summary.deals[0].company_name == "RawCo"


def test_direct_private_raw_company_folder_scan_uses_root_folder_name(tmp_path: Path) -> None:
    data_dir = tmp_path / "data"
    raw_company = data_dir / "raw" / "RawCo"
    closing = raw_company / "Closing"
    closing.mkdir(parents=True)
    (raw_company / "memo.txt").write_text("Memo about RawCo.", encoding="utf-8")
    (closing / "safe.txt").write_text("Simple Agreement for Future Equity.", encoding="utf-8")

    summary = ingest_folder(raw_company, config=AppConfig(data_dir=data_dir))

    assert summary.document_count == 2
    assert len(summary.deals) == 1
    assert summary.deals[0].company_name == "RawCo"
    assert {document.source.company_name for document in summary.deals[0].documents} == {"RawCo"}


def test_spreadsheets_are_recorded_with_extracted_text(tmp_path: Path) -> None:
    root = tmp_path / "pitch-decks"
    company = root / "SpreadsheetCo"
    company.mkdir(parents=True)
    (company / "revenue.csv").write_text("year,revenue\n2026,100\n", encoding="utf-8")
    with zipfile.ZipFile(company / "model.xlsx", "w") as workbook:
        workbook.writestr(
            "xl/sharedStrings.xml",
            """
            <sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
              <si><t>ARR</t></si>
            </sst>
            """,
        )
        workbook.writestr(
            "xl/worksheets/sheet1.xml",
            """
            <worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
              <sheetData>
                <row><c t="s"><v>0</v></c><c><v>250</v></c></row>
              </sheetData>
            </worksheet>
            """,
        )

    summary = ingest_folder(root, config=AppConfig(data_dir=tmp_path / "data"))

    assert summary.document_count == 2
    file_types = {doc.source.file_type for doc in summary.deals[0].documents}
    assert file_types == {FileType.CSV, FileType.XLSX}
    assert all(doc.pages for doc in summary.deals[0].documents)
    assert all(doc.source.table_count == 1 for doc in summary.deals[0].documents)
    assert all(doc.tables for doc in summary.deals[0].documents)

    saved_document = json.loads(summary.deals[0].documents[0].output_path.read_text())
    assert saved_document["source"]["table_count"] == 1
    assert saved_document["tables"]


def test_images_are_ingested_as_vision_needed_documents(tmp_path: Path) -> None:
    root = tmp_path / "pitch-decks"
    company = root / "ImageCo"
    company.mkdir(parents=True)
    (company / "scan.png").write_bytes(b"not real image bytes")
    (company / "photo.jpg").write_bytes(b"not real image bytes")

    summary = ingest_folder(root, config=AppConfig(data_dir=tmp_path / "data"))

    assert summary.document_count == 2
    assert summary.skipped_files == []
    assert summary.deals[0].evidence_count == 0
    documents = summary.deals[0].documents
    assert {document.source.file_type for document in documents} == {FileType.JPG, FileType.PNG}
    assert all(
        document.source.extraction_quality == ExtractionQuality.LOW
        for document in documents
    )
    assert all(document.source.ocr_recommended for document in documents)
    assert all(document.source.vision_recommended for document in documents)
    assert all(document.pages[0].needs_ocr for document in documents)
    assert all(document.pages[0].vision_recommended for document in documents)

    saved_document = json.loads(documents[0].output_path.read_text(encoding="utf-8"))
    assert saved_document["source"]["ocr_recommended"]
    assert saved_document["source"]["vision_recommended"]
    assert saved_document["pages"][0]["needs_ocr"]


@pytest.mark.skipif(not hasattr(os, "symlink"), reason="Symlinks are not supported here")
def test_output_directory_symlink_outside_data_dir_is_rejected(tmp_path: Path) -> None:
    root = tmp_path / "pitch-decks"
    company = root / "PrivateCo"
    company.mkdir(parents=True)
    (company / "memo.txt").write_text("Memo about PrivateCo.", encoding="utf-8")
    data_dir = tmp_path / "data"
    outside_dir = tmp_path / "outside"
    data_dir.mkdir()
    outside_dir.mkdir()
    (data_dir / "processed").symlink_to(outside_dir, target_is_directory=True)

    with pytest.raises(IngestionError, match="outside the private data directory"):
        ingest_folder(root, config=AppConfig(data_dir=data_dir))


@pytest.mark.skipif(not hasattr(os, "symlink"), reason="Symlinks are not supported here")
def test_final_summary_symlink_is_rejected_before_write(tmp_path: Path) -> None:
    root = tmp_path / "pitch-decks"
    company = root / "PrivateCo"
    company.mkdir(parents=True)
    (company / "memo.txt").write_text("Memo about PrivateCo.", encoding="utf-8")
    data_dir = tmp_path / "data"
    summary_dir = data_dir / "processed"
    summary_dir.mkdir(parents=True)
    outside_file = tmp_path / "outside-summary.json"
    outside_file.write_text("do not overwrite", encoding="utf-8")
    (summary_dir / "ingestion_summary.json").symlink_to(outside_file)

    with pytest.raises(IngestionError, match="output file is a symlink"):
        ingest_folder(root, config=AppConfig(data_dir=data_dir))

    assert outside_file.read_text(encoding="utf-8") == "do not overwrite"


def test_relative_data_dir_with_parent_segments_writes_outputs(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    workspace = tmp_path / "workspace"
    root = workspace / "pitch-decks"
    company = root / "RelativeCo"
    company.mkdir(parents=True)
    (company / "memo.txt").write_text("Memo about RelativeCo.", encoding="utf-8")
    monkeypatch.chdir(workspace)

    summary = ingest_folder(root, config=AppConfig(data_dir=Path("../hm-data")))

    assert summary.document_count == 1
    assert (tmp_path / "hm-data" / "processed" / "ingestion_summary.json").exists()
    assert summary.deals[0].documents[0].output_path.exists()


def test_generated_outputs_are_owner_only(tmp_path: Path) -> None:
    root = tmp_path / "pitch-decks"
    company = root / "PrivateCo"
    company.mkdir(parents=True)
    (company / "memo.txt").write_text("Memo about PrivateCo.", encoding="utf-8")

    summary = ingest_folder(root, config=AppConfig(data_dir=tmp_path / "data"))

    document_path = summary.deals[0].documents[0].output_path
    assert stat.S_IMODE(document_path.stat().st_mode) == 0o600
    assert stat.S_IMODE(summary.summary_path.stat().st_mode) == 0o600
    assert stat.S_IMODE((tmp_path / "data").stat().st_mode) == 0o700
    assert stat.S_IMODE((tmp_path / "data" / "processed").stat().st_mode) == 0o700
    assert stat.S_IMODE((tmp_path / "data" / "processed" / "deals").stat().st_mode) == 0o700
    assert stat.S_IMODE(document_path.parent.parent.stat().st_mode) == 0o700
    assert stat.S_IMODE(document_path.parent.stat().st_mode) == 0o700


def test_colliding_deal_folder_slugs_stay_separate(tmp_path: Path) -> None:
    root = tmp_path / "pitch-decks"
    first = root / "Acme.AI"
    second = root / "Acme AI"
    first.mkdir(parents=True)
    second.mkdir(parents=True)
    (first / "memo.txt").write_text("Memo about Acme.AI.", encoding="utf-8")
    (second / "memo.txt").write_text("Memo about Acme AI.", encoding="utf-8")

    summary = ingest_folder(root, config=AppConfig(data_dir=tmp_path / "data"))

    deal_ids = {deal.id for deal in summary.deals}
    company_names = {deal.company_name for deal in summary.deals}
    assert summary.document_count == 2
    assert len(summary.deals) == 2
    assert len(deal_ids) == 2
    assert company_names == {"Acme.AI", "Acme AI"}


def test_existing_deal_id_stays_stable_when_slug_collision_is_added(tmp_path: Path) -> None:
    root = tmp_path / "pitch-decks"
    original = root / "Acme.AI"
    original.mkdir(parents=True)
    (original / "memo.txt").write_text("Memo about Acme.AI.", encoding="utf-8")

    first_summary = ingest_folder(root, config=AppConfig(data_dir=tmp_path / "data"))
    first_deal_id = first_summary.deals[0].id

    new_collision = root / "Acme AI"
    new_collision.mkdir(parents=True)
    (new_collision / "memo.txt").write_text("Memo about Acme AI.", encoding="utf-8")

    second_summary = ingest_folder(root, config=AppConfig(data_dir=tmp_path / "data"))
    original_deal = next(
        deal for deal in second_summary.deals if deal.company_name == "Acme.AI"
    )

    assert original_deal.id == first_deal_id


def test_confidentiality_detection_uses_raw_text_before_cleanup(tmp_path: Path) -> None:
    root = tmp_path / "pitch-decks"
    company = root / "ConfidentialCo"
    company.mkdir(parents=True)
    (company / "ConfidentialCo memo.txt").write_text(
        "\n".join(
            [
                "Confidential: Parth Shah",
                "Business details.",
                "Confidential: Parth Shah",
                "More business details.",
                "Confidential: Parth Shah",
            ]
        ),
        encoding="utf-8",
    )

    summary = ingest_folder(root, config=AppConfig(data_dir=tmp_path / "data"))

    source = summary.deals[0].documents[0].source
    assert source.confidentiality_detected is True


def test_docx_header_confidentiality_is_detected(tmp_path: Path) -> None:
    root = tmp_path / "pitch-decks"
    company = root / "HeaderCo"
    company.mkdir(parents=True)
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Confidential"
    document.add_paragraph("Business details.")
    document.save(str(company / "memo.docx"))

    summary = ingest_folder(root, config=AppConfig(data_dir=tmp_path / "data"))

    source = summary.deals[0].documents[0].source
    assert source.confidentiality_detected is True


def test_confidentiality_detection_scans_all_raw_text(tmp_path: Path) -> None:
    root = tmp_path / "pitch-decks"
    company = root / "LongCo"
    company.mkdir(parents=True)
    (company / "long memo.txt").write_text(
        f"{'A' * 21_000}\nNot for distribution.",
        encoding="utf-8",
    )

    summary = ingest_folder(root, config=AppConfig(data_dir=tmp_path / "data"))

    source = summary.deals[0].documents[0].source
    assert source.confidentiality_detected is True


def test_confidentiality_detection_uses_filename_when_text_is_empty(tmp_path: Path) -> None:
    root = tmp_path / "pitch-decks"
    company = root / "LockedCo"
    company.mkdir(parents=True)
    (company / "Confidential Investor Deck.pdf").write_text("not a real pdf", encoding="utf-8")

    summary = ingest_folder(root, config=AppConfig(data_dir=tmp_path / "data"))

    source = summary.deals[0].documents[0].source
    assert source.confidentiality_detected is True


def test_hash_failure_is_recorded_without_aborting(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = tmp_path / "pitch-decks"
    company = root / "HashFailCo"
    company.mkdir(parents=True)
    (company / "memo.txt").write_text("Memo about HashFailCo.", encoding="utf-8")

    def fail_sha256(path: Path) -> str:
        raise PermissionError("permission changed")

    monkeypatch.setattr(folder_loader, "_sha256", fail_sha256)

    summary = ingest_folder(root, config=AppConfig(data_dir=tmp_path / "data"))

    source = summary.deals[0].documents[0].source
    assert summary.document_count == 1
    assert source.sha256 is None
    assert source.notes is not None
    assert "Could not calculate the file fingerprint" in source.notes
    assert summary.deals[0].documents[0].output_path.exists()


@pytest.mark.skipif(not hasattr(os, "symlink"), reason="Symlinks are not supported here")
def test_symlinked_files_are_skipped(tmp_path: Path) -> None:
    root = tmp_path / "pitch-decks"
    company = root / "SymlinkCo"
    company.mkdir(parents=True)
    outside_file = tmp_path / "outside-secret.txt"
    outside_file.write_text("This should not be read.", encoding="utf-8")
    (company / "deck.txt").symlink_to(outside_file)

    summary = ingest_folder(root, config=AppConfig(data_dir=tmp_path / "data"))

    assert summary.document_count == 0
    assert summary.skipped_files == ["SymlinkCo/deck.txt"]


@pytest.mark.skipif(not hasattr(os, "symlink"), reason="Symlinks are not supported here")
def test_symlinked_scan_root_is_rejected_before_resolving(tmp_path: Path) -> None:
    real_root = tmp_path / "outside-decks"
    company = real_root / "SecretCo"
    company.mkdir(parents=True)
    (company / "memo.txt").write_text("This should not be scanned.", encoding="utf-8")
    symlinked_root = tmp_path / "pitch-decks"
    symlinked_root.symlink_to(real_root, target_is_directory=True)

    with pytest.raises(IngestionError, match="scan folder cannot be a symlink"):
        ingest_folder(symlinked_root, config=AppConfig(data_dir=tmp_path / "data"))

    assert not (tmp_path / "data" / "processed" / "ingestion_summary.json").exists()


def test_unreadable_scan_folders_are_reported(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = tmp_path / "pitch-decks"
    company = root / "HiddenCo"
    secret_dir = company / "Secret"
    company.mkdir(parents=True)
    secret_dir.mkdir()
    (company / "memo.txt").write_text("Memo about HiddenCo.", encoding="utf-8")

    def fake_walk(
        top: Path,
        topdown: bool,
        onerror: Callable[[OSError], None] | None,
        followlinks: bool,
    ) -> Iterator[tuple[Path, list[str], list[str]]]:
        assert top == root
        assert topdown is True
        assert followlinks is False
        yield root, ["HiddenCo"], []
        if callable(onerror):
            onerror(PermissionError(13, "Permission denied", str(secret_dir)))
        yield company, ["Secret"], ["memo.txt"]

    monkeypatch.setattr(os, "walk", fake_walk)

    summary = ingest_folder(root, config=AppConfig(data_dir=tmp_path / "data"))

    assert summary.document_count == 1
    assert summary.skipped_files == []
    assert summary.unreadable_paths == ["HiddenCo/Secret"]
