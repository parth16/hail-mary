from __future__ import annotations

import csv
import io
import re
import zipfile
from collections.abc import Iterable
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from bs4 import BeautifulSoup
from docx import Document
from pydantic import BaseModel, Field
from pypdf import PdfReader

from hailmary.ingest.document_classifier import classify_file_type
from hailmary.ingest.ocr import (
    LOW_OCR_CONFIDENCE_THRESHOLD,
    LocalOcrDependencyError,
    LocalOcrEngine,
    LocalOcrError,
    LocalOcrResult,
)
from hailmary.schemas.documents import (
    ExtractedPage,
    ExtractedTable,
    ExtractionQuality,
    FileType,
)
from hailmary.utils.text_cleaning import clean_extracted_text_with_metadata

MAX_XLSX_COLUMN_INDEX = 16_383
MAX_XLSX_COLUMN_LETTERS = 3
MAX_XLSX_BLANK_GAP = 100
INVALID_XLSX_CELL_INDEX = -1
PDF_REPEATED_SHORT_TEXT_MIN_PAGES = 2
PDF_REPEATED_SHORT_TEXT_WORD_THRESHOLD = 8
PDF_IMAGE_TEXT_WORD_THRESHOLD = 8
MAX_HTML_TABLE_SPAN = 100
LOCAL_OCR_DOCUMENT_NOTE = (
    "Some pages may need local OCR, image-based text reading (OCR), before Hail Mary "
    "can use all of their content. OCR means reading text from images."
)
LOCAL_OCR_EMPTY_PAGE_NOTE = (
    "Page has no extracted text and may need image-based text reading (OCR). "
    "OCR means reading text from images."
)
LOCAL_OCR_IMAGE_PAGE_NOTE = (
    "Page has image content with little extracted text and may need image-based text "
    "reading (OCR). OCR means reading text from images."
)
IMAGE_LOCAL_OCR_NOTE = (
    "Image file needs local OCR, image-based text reading (OCR), before text can be "
    "extracted. OCR means reading text from images."
)
OCR_APPLIED_NOTE = "Image-based text reading (OCR) was used for this text."
OCR_ATTEMPTED_NOTE = (
    "Image-based text reading (OCR) was attempted. OCR means reading text from images."
)
OCR_LOW_TEXT_NOTE = (
    "Image-based text reading (OCR) found only a small amount of text. "
    "Review the source image because most content may still be unreadable."
)
OCR_PRESERVED_TEXT_NOTE = (
    "Existing PDF text was preserved while image-based text reading (OCR) was used."
)


class ExtractionResult(BaseModel):
    pages: list[ExtractedPage]
    tables: list[ExtractedTable] = Field(default_factory=list)
    page_count: int | None
    extraction_quality: ExtractionQuality
    ocr_recommended: bool = False
    ocr_applied: bool = False
    ocr_confidence: float | None = Field(default=None, ge=0.0, le=1.0)
    vision_recommended: bool = False
    notes: str | None = None

    @property
    def combined_text(self) -> str:
        return "\n\n".join(page.clean_text for page in self.pages if page.clean_text)

    @property
    def combined_raw_text(self) -> str:
        return "\n\n".join(page.raw_text for page in self.pages if page.raw_text)

    @property
    def table_count(self) -> int:
        return len(self.tables)


def extract_document(
    path: Path,
    *,
    ocr_engine: LocalOcrEngine | None = None,
) -> ExtractionResult:
    file_type = classify_file_type(path)

    if file_type == FileType.PDF:
        return _extract_pdf(path, ocr_engine=ocr_engine)
    if file_type == FileType.DOCX:
        return _extract_docx(path)
    if file_type == FileType.HTML:
        return _extract_html(path)
    if file_type == FileType.CSV:
        return _extract_csv(path)
    if file_type == FileType.XLSX:
        return _extract_xlsx(path)
    if file_type in {FileType.TXT, FileType.MD}:
        return _extract_text_file(path)
    if file_type in {FileType.PNG, FileType.JPG}:
        return _extract_image(path, ocr_engine=ocr_engine)

    return ExtractionResult(
        pages=[],
        page_count=None,
        extraction_quality=ExtractionQuality.LOW,
        notes="Text extraction is not implemented for this file type yet.",
    )


def _quality_from_pages(pages: list[ExtractedPage]) -> ExtractionQuality:
    if not pages:
        return ExtractionQuality.LOW

    word_count = sum(page.word_count for page in pages)
    words_per_page = word_count / len(pages)
    ocr_needed_pages = sum(1 for page in pages if page.needs_ocr)
    if words_per_page >= 50 and ocr_needed_pages * 2 < len(pages):
        return ExtractionQuality.HIGH
    if word_count > 0:
        return ExtractionQuality.MEDIUM
    return ExtractionQuality.LOW


def _make_page(
    raw_text: str,
    *,
    page_number: int | None,
    needs_ocr: bool | None,
    vision_recommended: bool | None = False,
    source_span_start: int | None = None,
    notes: str | None = None,
) -> ExtractedPage:
    cleaning = clean_extracted_text_with_metadata(raw_text)
    word_count = len(cleaning.clean_text.split())
    page_needs_ocr = (
        _clean_text_needs_ocr(cleaning.clean_text, word_count=word_count)
        if needs_ocr is None
        else needs_ocr
    )
    page_needs_vision = page_needs_ocr if vision_recommended is None else vision_recommended
    source_span_end = (
        source_span_start + len(raw_text) if source_span_start is not None else None
    )
    if needs_ocr is None and page_needs_ocr and not raw_text.strip():
        notes = _append_note(notes, LOCAL_OCR_EMPTY_PAGE_NOTE)
    return ExtractedPage(
        page_number=page_number,
        raw_text=raw_text,
        clean_text=cleaning.clean_text,
        word_count=word_count,
        needs_ocr=page_needs_ocr,
        vision_recommended=page_needs_vision,
        source_span_start=source_span_start,
        source_span_end=source_span_end,
        removed_boilerplate_lines=cleaning.removed_boilerplate_lines,
        notes=notes,
    )


def _append_note(existing_note: str | None, new_note: str) -> str:
    if not existing_note:
        return new_note
    if new_note in existing_note:
        return existing_note
    return f"{existing_note} {new_note}"


def _clean_text_needs_ocr(clean_text: str, *, word_count: int) -> bool:
    if not clean_text.strip():
        return True
    return word_count <= 2


def _mark_repeated_short_pdf_pages_for_ocr(pages: list[ExtractedPage]) -> None:
    short_text_pages = [
        page
        for page in pages
        if page.raw_text.strip()
        and not page.notes
        and 2 < page.word_count <= PDF_REPEATED_SHORT_TEXT_WORD_THRESHOLD
    ]
    if len(short_text_pages) < PDF_REPEATED_SHORT_TEXT_MIN_PAGES:
        return
    if len(short_text_pages) * 2 <= len(pages):
        return

    for page in short_text_pages:
        page.needs_ocr = True
        page.vision_recommended = True


def _result_from_pages(
    pages: list[ExtractedPage],
    *,
    page_count: int | None,
    tables: list[ExtractedTable] | None = None,
    notes: str | None = None,
) -> ExtractionResult:
    ocr_recommended = _document_ocr_recommended(pages)
    ocr_confidence = _ocr_confidence_from_pages(pages)
    result_notes = notes
    if ocr_recommended:
        result_notes = _append_note(result_notes, LOCAL_OCR_DOCUMENT_NOTE)
    return ExtractionResult(
        pages=pages,
        tables=tables or [],
        page_count=page_count,
        extraction_quality=_quality_from_pages(pages),
        ocr_recommended=ocr_recommended,
        ocr_applied=any(page.ocr_applied for page in pages),
        ocr_confidence=ocr_confidence,
        vision_recommended=ocr_recommended
        or any(page.vision_recommended and page.notes for page in pages),
        notes=result_notes,
    )


def _ocr_confidence_from_pages(pages: list[ExtractedPage]) -> float | None:
    confidences = [
        page.ocr_confidence
        for page in pages
        if page.ocr_applied and page.ocr_confidence is not None
    ]
    if not confidences:
        return None
    return sum(confidences) / len(confidences)


def _document_ocr_recommended(pages: list[ExtractedPage]) -> bool:
    if not pages:
        return False

    ocr_pages = [page for page in pages if page.needs_ocr]
    if not ocr_pages:
        return False
    if any(page.notes for page in ocr_pages):
        return True
    if any(not page.raw_text.strip() for page in ocr_pages):
        return True

    return len(ocr_pages) * 2 > len(pages)


def _extract_pdf(
    path: Path,
    *,
    ocr_engine: LocalOcrEngine | None,
) -> ExtractionResult:
    try:
        reader = PdfReader(path)
    except Exception as exc:
        return ExtractionResult(
            pages=[],
            page_count=None,
            extraction_quality=ExtractionQuality.LOW,
            notes=f"Could not read the PDF: {exc}",
        )

    try:
        pages_proxy = reader.pages
        page_count = len(pages_proxy)
    except Exception as exc:
        return ExtractionResult(
            pages=[],
            page_count=None,
            extraction_quality=ExtractionQuality.LOW,
            notes=f"Could not read pages from the PDF: {exc}",
        )

    pages: list[ExtractedPage] = []
    source_offset = 0
    for index in range(page_count):
        has_image_resources = False
        try:
            page = pages_proxy[index]
            raw_text = page.extract_text() or ""
            notes = None
            has_image_resources = _pdf_page_has_image_resources(page)
        except Exception as exc:
            raw_text = ""
            notes = f"Could not extract text from page {index + 1}: {exc}"

        extracted_page = _make_page(
            raw_text,
            page_number=index + 1,
            needs_ocr=None,
            vision_recommended=None,
            source_span_start=source_offset,
            notes=notes,
        )
        if has_image_resources and extracted_page.word_count <= PDF_IMAGE_TEXT_WORD_THRESHOLD:
            extracted_page.needs_ocr = True
            extracted_page.vision_recommended = True
            extracted_page.notes = _append_note(
                extracted_page.notes,
                LOCAL_OCR_IMAGE_PAGE_NOTE,
            )
        pages.append(extracted_page)
        if raw_text:
            source_offset += len(raw_text) + 2

    _mark_repeated_short_pdf_pages_for_ocr(pages)
    if ocr_engine is not None:
        _apply_pdf_ocr(path, pages, ocr_engine)
        _refresh_page_source_spans(pages)
    return _result_from_pages(
        pages,
        page_count=page_count,
        notes=_page_failure_notes(pages),
    )


def _apply_pdf_ocr(
    path: Path,
    pages: list[ExtractedPage],
    ocr_engine: LocalOcrEngine,
) -> None:
    targets = _pdf_ocr_target_pages(pages)
    for page in targets:
        if page.page_number is None:
            continue
        try:
            ocr_result = ocr_engine.pdf_page_to_text(path, page_number=page.page_number)
        except LocalOcrDependencyError as exc:
            page.notes = _append_note(page.notes, str(exc))
            return
        except LocalOcrError as exc:
            page.notes = _append_note(page.notes, str(exc))
            continue
        _apply_ocr_result_to_page(page, ocr_result)


def _pdf_ocr_target_pages(pages: list[ExtractedPage]) -> list[ExtractedPage]:
    document_needs_ocr = _document_ocr_recommended(pages)
    return [
        page
        for page in pages
        if page.needs_ocr
        and (
            document_needs_ocr
            or not page.raw_text.strip()
            or _page_note_contains(page, LOCAL_OCR_IMAGE_PAGE_NOTE)
        )
    ]


def _page_note_contains(page: ExtractedPage, note: str) -> bool:
    return page.notes is not None and note in page.notes


def _pdf_page_has_image_resources(page: Any) -> bool:
    if _pdf_page_images_collection_has_items(page):
        return True

    try:
        resources = _pdf_resolved_object(page.get("/Resources"))
        if not resources:
            return False
        xobjects = _pdf_resolved_object(resources.get("/XObject"))
        if not xobjects:
            return False
        for xobject in xobjects.values():
            resolved_xobject = _pdf_resolved_object(xobject)
            xobject_get = getattr(resolved_xobject, "get", None)
            if xobject_get is not None and str(xobject_get("/Subtype")) == "/Image":
                return True
    except Exception:
        return False
    return False


def _pdf_page_images_collection_has_items(page: Any) -> bool:
    try:
        images = getattr(page, "images", None)
        if images is None:
            return False
        try:
            return len(images) > 0
        except TypeError:
            return any(True for _ in images)
    except Exception:
        return False


def _pdf_resolved_object(value: Any) -> Any:
    get_object = getattr(value, "get_object", None)
    if get_object is None:
        return value
    return get_object()


def _extract_image(
    path: Path,
    *,
    ocr_engine: LocalOcrEngine | None,
) -> ExtractionResult:
    try:
        path.stat()
    except OSError as exc:
        return ExtractionResult(
            pages=[],
            page_count=None,
            extraction_quality=ExtractionQuality.LOW,
            notes=f"Could not read the image file: {exc}",
        )

    if ocr_engine is not None:
        try:
            ocr_result = ocr_engine.image_to_text(path, page_number=1)
        except LocalOcrError as exc:
            return _image_needs_ocr_result(str(exc))
        page = ExtractedPage(
            page_number=1,
            raw_text="",
            clean_text="",
            word_count=0,
            needs_ocr=True,
            vision_recommended=True,
            source_span_start=0,
            source_span_end=0,
        )
        _apply_ocr_result_to_page(page, ocr_result)
        if not ocr_result.text.strip():
            page.notes = _append_note(page.notes, IMAGE_LOCAL_OCR_NOTE)
        return _result_from_pages(
            [page],
            page_count=1,
            notes=page.notes,
        )

    return _image_needs_ocr_result(IMAGE_LOCAL_OCR_NOTE)


def _image_needs_ocr_result(notes: str) -> ExtractionResult:
    page = ExtractedPage(
        page_number=1,
        raw_text="",
        clean_text="",
        word_count=0,
        needs_ocr=True,
        vision_recommended=True,
        source_span_start=0,
        source_span_end=0,
        notes=notes,
    )
    return ExtractionResult(
        pages=[page],
        page_count=1,
        extraction_quality=ExtractionQuality.LOW,
        ocr_recommended=True,
        vision_recommended=True,
        notes=notes,
    )


def _apply_ocr_result_to_page(
    page: ExtractedPage,
    ocr_result: LocalOcrResult,
) -> None:
    if not ocr_result.text.strip():
        page.notes = _append_note(
            page.notes,
            _ocr_page_notes(ocr_result, applied=False),
        )
        return

    if _ocr_result_is_low_confidence(ocr_result):
        page.ocr_confidence = ocr_result.confidence
        if page.clean_text.strip():
            page.needs_ocr = True
            page.vision_recommended = True
            page.notes = _append_note(
                page.notes,
                _ocr_page_notes(ocr_result, applied=False),
            )
            return

        page.ocr_applied = True
        updated_page = _make_page(
            ocr_result.text,
            page_number=page.page_number,
            needs_ocr=True,
            vision_recommended=True,
            source_span_start=page.source_span_start,
            notes=_ocr_page_notes(ocr_result),
        )
        page.raw_text = updated_page.raw_text
        page.clean_text = ""
        page.word_count = 0
        page.needs_ocr = True
        page.vision_recommended = True
        page.source_span_end = updated_page.source_span_end
        page.removed_boilerplate_lines = updated_page.removed_boilerplate_lines
        page.notes = updated_page.notes
        return

    page.ocr_confidence = ocr_result.confidence
    updated_page = _make_page(
        ocr_result.text,
        page_number=page.page_number,
        needs_ocr=None,
        vision_recommended=None,
        source_span_start=page.source_span_start,
        notes=_ocr_page_notes(ocr_result),
    )
    clean_text = updated_page.clean_text
    word_count = updated_page.word_count
    notes = updated_page.notes
    ocr_text_is_insufficient = updated_page.needs_ocr
    if updated_page.needs_ocr:
        if page.clean_text.strip():
            page.needs_ocr = True
            page.vision_recommended = True
            page.notes = _append_note(
                page.notes,
                _ocr_page_notes(ocr_result, applied=False),
            )
            page.notes = _append_note(page.notes, OCR_LOW_TEXT_NOTE)
            return

        clean_text = ""
        word_count = 0
        notes = _append_note(notes, OCR_LOW_TEXT_NOTE)
    merged_raw_text = _merge_existing_page_text_with_ocr(
        existing_text=page.raw_text,
        ocr_text=updated_page.raw_text,
    )
    if merged_raw_text != updated_page.raw_text:
        updated_page = _make_page(
            merged_raw_text,
            page_number=page.page_number,
            needs_ocr=updated_page.needs_ocr,
            vision_recommended=updated_page.vision_recommended,
            source_span_start=page.source_span_start,
            notes=_append_note(notes, OCR_PRESERVED_TEXT_NOTE),
        )
        if ocr_text_is_insufficient:
            clean_text = ""
            word_count = 0
        else:
            clean_text = updated_page.clean_text
            word_count = updated_page.word_count
        notes = updated_page.notes
    page.ocr_applied = True
    page.raw_text = updated_page.raw_text
    page.clean_text = clean_text
    page.word_count = word_count
    page.needs_ocr = updated_page.needs_ocr
    page.vision_recommended = updated_page.vision_recommended
    page.source_span_end = updated_page.source_span_end
    page.removed_boilerplate_lines = updated_page.removed_boilerplate_lines
    page.notes = notes


def _ocr_result_is_low_confidence(ocr_result: LocalOcrResult) -> bool:
    return (
        ocr_result.confidence is not None
        and ocr_result.confidence < LOW_OCR_CONFIDENCE_THRESHOLD
    )


def _ocr_page_notes(ocr_result: LocalOcrResult, *, applied: bool = True) -> str:
    notes = _append_note(None, OCR_APPLIED_NOTE if applied else OCR_ATTEMPTED_NOTE)
    if _ocr_result_is_low_confidence(ocr_result):
        notes = _append_note(
            notes,
            "Image-based text reading (OCR) finished with low confidence. "
            "Review the source image before relying on this text.",
        )
    if ocr_result.notes:
        notes = _append_note(notes, ocr_result.notes)
    return notes


def _merge_existing_page_text_with_ocr(*, existing_text: str, ocr_text: str) -> str:
    existing = existing_text.strip()
    ocr = ocr_text.strip()
    if not existing:
        return ocr
    if not ocr:
        return existing

    existing_lines = _normalized_nonblank_lines(existing)
    ocr_lines = _normalized_nonblank_lines(ocr)
    if existing_lines == ocr_lines:
        return existing
    existing_line_set = set(existing_lines)
    ocr_line_set = set(ocr_lines)
    if existing_lines and set(existing_lines).issubset(ocr_line_set):
        return ocr
    if ocr_lines and set(ocr_lines).issubset(existing_line_set):
        return existing

    merged_lines = [line.strip() for line in existing.splitlines() if line.strip()]
    merged_lines.extend(
        line.strip()
        for line in ocr.splitlines()
        if line.strip() and _normalize_text_for_merge(line) not in existing_line_set
    )
    return "\n\n".join(merged_lines)


def _normalized_nonblank_lines(text: str) -> list[str]:
    return [
        normalized_line
        for line in text.splitlines()
        if (normalized_line := _normalize_text_for_merge(line))
    ]


def _normalize_text_for_merge(text: str) -> str:
    return re.sub(r"\s+", " ", text.casefold()).strip()


def _refresh_page_source_spans(pages: list[ExtractedPage]) -> None:
    source_offset = 0
    for page in pages:
        page.source_span_start = source_offset
        page.source_span_end = source_offset + len(page.raw_text)
        if page.raw_text:
            source_offset += len(page.raw_text) + 2


def _extract_docx(path: Path) -> ExtractionResult:
    try:
        document = Document(str(path))
    except Exception as exc:
        return ExtractionResult(
            pages=[],
            page_count=None,
            extraction_quality=ExtractionQuality.LOW,
            notes=f"Could not read the DOCX file: {exc}",
        )

    try:
        text_parts = _docx_text_parts(document)
        tables = _docx_tables(document)
    except Exception as exc:
        return ExtractionResult(
            pages=[],
            page_count=None,
            extraction_quality=ExtractionQuality.LOW,
            notes=f"Could not extract text from the DOCX file: {exc}",
        )

    table_text = [table.clean_text for table in tables if table.clean_text]
    raw_text = "\n".join([*text_parts, *table_text])
    page = _make_page(
        raw_text,
        page_number=None,
        needs_ocr=False,
        source_span_start=0,
    )

    pages = _pages_with_raw_text(page)
    return _result_from_pages(
        pages,
        tables=tables,
        page_count=None,
    )


def _docx_text_parts(document: Any) -> list[str]:
    text_parts: list[str] = []
    text_parts.extend(_paragraph_text(document.paragraphs))

    for section in document.sections:
        containers = [
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ]
        for container in containers:
            text_parts.extend(_paragraph_text(container.paragraphs))

    return text_parts


def _paragraph_text(paragraphs: Iterable[Any]) -> list[str]:
    return [
        paragraph.text.strip()
        for paragraph in paragraphs
        if getattr(paragraph, "text", "").strip()
    ]


def _docx_tables(document: Any) -> list[ExtractedTable]:
    tables: list[ExtractedTable] = []
    table_index = 1
    for table in _iter_docx_tables(document):
        rows = _docx_table_rows(table)
        if not rows:
            continue
        tables.append(_make_table(rows, table_index=table_index, page_number=None))
        table_index += 1
    return tables


def _iter_docx_tables(document: Any) -> Iterable[Any]:
    yield from document.tables
    for section in document.sections:
        containers = [
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ]
        for container in containers:
            yield from container.tables


def _docx_table_rows(table: Any) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in table.rows:
        cells = [" ".join(cell.text.split()) for cell in row.cells]
        while cells and not cells[-1]:
            cells.pop()
        if any(cells):
            rows.append(cells)
    return rows


def _page_failure_notes(pages: list[ExtractedPage]) -> str | None:
    page_notes = [page.notes for page in pages if page.notes]
    if not page_notes:
        return None
    return " ".join(page_notes)


def _make_table(
    rows: list[list[str]],
    *,
    table_index: int,
    page_number: int | None = None,
    source_span_start: int | None = None,
    notes: str | None = None,
    preserve_trailing_blanks: bool = False,
) -> ExtractedTable:
    normalized_rows = _table_rows_with_content(
        rows,
        preserve_trailing_blanks=preserve_trailing_blanks,
    )
    text = "\n".join(_table_rows_as_text(normalized_rows))
    source_span_end = source_span_start + len(text) if source_span_start is not None else None
    return ExtractedTable(
        page_number=page_number,
        table_index=table_index,
        rows=normalized_rows,
        clean_text=text,
        row_count=len(normalized_rows),
        column_count=max((len(row) for row in normalized_rows), default=0),
        source_span_start=source_span_start,
        source_span_end=source_span_end,
        notes=notes,
    )


def _table_rows_as_text(rows: list[list[str]]) -> list[str]:
    return [" | ".join(row) for row in _table_rows_with_content(rows)]


def _table_rows_with_content(
    rows: list[list[str]],
    *,
    preserve_trailing_blanks: bool = False,
) -> list[list[str]]:
    normalized_rows: list[list[str]] = []
    for row in rows:
        cells = [cell.strip() for cell in row]
        if not preserve_trailing_blanks:
            cells = _trim_trailing_blank_cells(cells)
        if any(cells):
            normalized_rows.append(cells)
    return normalized_rows


def _trim_trailing_blank_cells(row: list[str]) -> list[str]:
    cells = list(row)
    while cells and not cells[-1]:
        cells.pop()
    return cells


def _extract_text_file(path: Path) -> ExtractionResult:
    try:
        raw_text = path.read_text(encoding="utf-8")
        notes = None
    except UnicodeDecodeError:
        raw_text = path.read_text(encoding="utf-8", errors="replace")
        notes = "Some characters could not be read and were replaced."
    except OSError as exc:
        return ExtractionResult(
            pages=[],
            page_count=None,
            extraction_quality=ExtractionQuality.LOW,
            notes=f"Could not read the text file: {exc}",
        )

    page = _make_page(
        raw_text,
        page_number=None,
        needs_ocr=False,
        source_span_start=0,
        notes=notes,
    )

    pages = _pages_with_raw_text(page)
    return _result_from_pages(
        pages,
        page_count=1,
        notes=notes,
    )


def _extract_html(path: Path) -> ExtractionResult:
    try:
        html = path.read_text(encoding="utf-8")
        notes = None
    except UnicodeDecodeError:
        html = path.read_text(encoding="utf-8", errors="replace")
        notes = "Some characters could not be read and were replaced."
    except OSError as exc:
        return ExtractionResult(
            pages=[],
            page_count=None,
            extraction_quality=ExtractionQuality.LOW,
            notes=f"Could not read the HTML file: {exc}",
        )

    soup = BeautifulSoup(html, "html.parser")
    for element in soup(["script", "style", "noscript"]):
        element.decompose()

    tables = _html_tables(soup)
    raw_text = soup.get_text(separator="\n")
    page = _make_page(
        raw_text,
        page_number=None,
        needs_ocr=False,
        source_span_start=0,
        notes=notes,
    )

    pages = _pages_with_raw_text(page)
    return _result_from_pages(
        pages,
        tables=tables,
        page_count=1,
        notes=notes,
    )


def _html_tables(soup: BeautifulSoup) -> list[ExtractedTable]:
    tables: list[ExtractedTable] = []
    for table_index, table in enumerate(soup.find_all("table"), start=1):
        rows = _html_table_rows(table)
        if rows:
            tables.append(
                _make_table(
                    rows,
                    table_index=table_index,
                    preserve_trailing_blanks=True,
                )
            )
    return tables


def _html_table_rows(table: Any) -> list[list[str]]:
    rows: list[list[str]] = []
    active_rowspans: dict[int, tuple[int, str]] = {}
    for row in _html_direct_table_rows(table):
        cells: list[str] = []
        column_index = 0
        row_has_span = bool(active_rowspans)
        for cell in row.find_all(["th", "td"], recursive=False):
            column_index = _append_html_rowspans(
                cells, active_rowspans, column_index
            )
            cell_text = _html_cell_text(cell)
            rowspan = _html_span_value(cell, "rowspan")
            colspan = _html_span_value(cell, "colspan")
            row_has_span = row_has_span or rowspan > 1 or colspan > 1
            cells.append(cell_text)
            cells.extend([""] * (colspan - 1))
            if rowspan > 1:
                for span_offset in range(colspan):
                    active_rowspans[column_index + span_offset] = (
                        rowspan - 1,
                        cell_text if span_offset == 0 else "",
                    )
            column_index += colspan

        _append_html_rowspans(cells, active_rowspans, column_index)
        while not row_has_span and cells and not cells[-1]:
            cells.pop()
        if any(cells):
            rows.append(cells)
    return rows


def _append_html_rowspans(
    cells: list[str],
    active_rowspans: dict[int, tuple[int, str]],
    column_index: int,
) -> int:
    while column_index in active_rowspans:
        remaining_rows, cell_text = active_rowspans[column_index]
        cells.append(cell_text)
        if remaining_rows <= 1:
            del active_rowspans[column_index]
        else:
            active_rowspans[column_index] = (remaining_rows - 1, cell_text)
        column_index += 1
    return column_index


def _html_span_value(cell: Any, attribute: str) -> int:
    raw_value = cell.get(attribute)
    if raw_value is None:
        return 1
    try:
        value = int(raw_value)
    except (TypeError, ValueError):
        return 1
    return min(max(value, 1), MAX_HTML_TABLE_SPAN)


def _html_direct_table_rows(table: Any) -> list[Any]:
    rows: list[Any] = []
    for child in table.find_all(["thead", "tbody", "tfoot", "tr"], recursive=False):
        if child.name == "tr":
            rows.append(child)
            continue
        rows.extend(child.find_all("tr", recursive=False))
    return rows


def _html_cell_text(cell: Any) -> str:
    copied_cell_soup = BeautifulSoup(str(cell), "html.parser")
    copied_cell = copied_cell_soup.find(["th", "td"])
    if copied_cell is None:
        return ""
    for nested_table in copied_cell.find_all("table"):
        nested_table.decompose()
    return " ".join(copied_cell.get_text(" ").split())


def _extract_csv(path: Path) -> ExtractionResult:
    try:
        csv_text = path.read_text(encoding="utf-8")
        notes = None
    except UnicodeDecodeError:
        csv_text = path.read_text(encoding="utf-8", errors="replace")
        notes = "Some characters could not be read and were replaced."
    except OSError as exc:
        return ExtractionResult(
            pages=[],
            page_count=None,
            extraction_quality=ExtractionQuality.LOW,
            notes=f"Could not read the CSV file: {exc}",
        )

    try:
        rows = [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(csv_text))]
    except csv.Error as exc:
        return ExtractionResult(
            pages=[],
            page_count=None,
            extraction_quality=ExtractionQuality.LOW,
            notes=f"Could not parse the CSV file: {exc}",
        )
    content_rows = _table_rows_with_content(rows)
    raw_text = "\n".join(_table_rows_as_text(content_rows))
    tables = [_make_table(content_rows, table_index=1, source_span_start=0)] if content_rows else []
    return _single_page_result(raw_text, page_count=1, notes=notes, tables=tables)


def _extract_xlsx(path: Path) -> ExtractionResult:
    try:
        with zipfile.ZipFile(path) as workbook:
            shared_strings = _xlsx_shared_strings(workbook)
            sheet_names = sorted(
                name
                for name in workbook.namelist()
                if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
            )
            sheet_rows = [
                _xlsx_sheet_rows(workbook.read(sheet_name), shared_strings)
                for sheet_name in sheet_names
            ]
    except (OSError, RuntimeError, zipfile.BadZipFile, ElementTree.ParseError) as exc:
        return ExtractionResult(
            pages=[],
            page_count=None,
            extraction_quality=ExtractionQuality.LOW,
            notes=f"Could not read the XLSX file: {exc}",
        )

    tables = [
        _make_table(rows, table_index=index)
        for index, rows in enumerate(sheet_rows, start=1)
        if rows
    ]
    raw_text = "\n".join(table.clean_text for table in tables if table.clean_text)
    return _single_page_result(
        raw_text,
        page_count=len(sheet_names),
        notes=None,
        tables=tables,
    )


def _xlsx_shared_strings(workbook: zipfile.ZipFile) -> list[str]:
    try:
        shared_strings_xml = workbook.read("xl/sharedStrings.xml")
    except KeyError:
        return []

    root = ElementTree.fromstring(shared_strings_xml)
    strings: list[str] = []
    for item in root.iter():
        if _xml_local_name(item.tag) != "si":
            continue
        text_parts = [
            text_element.text or ""
            for text_element in item.iter()
            if _xml_local_name(text_element.tag) == "t"
        ]
        strings.append("".join(text_parts).strip())
    return strings


def _xlsx_sheet_rows(sheet_xml: bytes, shared_strings: list[str]) -> list[list[str]]:
    root = ElementTree.fromstring(sheet_xml)
    rows: list[list[str]] = []
    for row in root.iter():
        if _xml_local_name(row.tag) != "row":
            continue
        values: list[str] = []
        next_source_column = 0
        for cell in row:
            if _xml_local_name(cell.tag) != "c":
                continue
            cell_index = _xlsx_cell_index(cell)
            if cell_index == INVALID_XLSX_CELL_INDEX:
                continue
            cell_value = _xlsx_cell_value(cell, shared_strings)
            if not cell_value:
                continue
            if cell_index is not None:
                if cell_index > MAX_XLSX_COLUMN_INDEX:
                    continue
                gap = cell_index - next_source_column
                if gap > MAX_XLSX_BLANK_GAP:
                    values.append(f"[{gap} blank columns]")
                else:
                    values.extend([""] * max(gap, 0))
                next_source_column = max(next_source_column, cell_index + 1)
            else:
                next_source_column += 1
            values.append(cell_value)
        while values and not values[-1]:
            values.pop()
        if any(values):
            rows.append(values)
    return rows


def _xlsx_cell_index(cell: ElementTree.Element) -> int | None:
    cell_reference = cell.attrib.get("r")
    if not cell_reference:
        return None
    match = re.match(r"([A-Za-z]+)", cell_reference)
    if not match:
        return INVALID_XLSX_CELL_INDEX

    column_name = match.group(1).upper()
    if len(column_name) > MAX_XLSX_COLUMN_LETTERS:
        return INVALID_XLSX_CELL_INDEX

    index = 0
    for character in column_name:
        index = index * 26 + (ord(character) - ord("A") + 1)
    zero_based_index = index - 1
    if zero_based_index > MAX_XLSX_COLUMN_INDEX:
        return INVALID_XLSX_CELL_INDEX
    return zero_based_index


def _xlsx_cell_value(cell: ElementTree.Element, shared_strings: list[str]) -> str:
    cell_type = cell.attrib.get("t")
    if cell_type == "inlineStr":
        return " ".join(
            text_element.text or ""
            for text_element in cell.iter()
            if _xml_local_name(text_element.tag) == "t"
        ).strip()

    value = ""
    for child in cell:
        if _xml_local_name(child.tag) == "v":
            value = child.text or ""
            break

    if cell_type == "s":
        try:
            shared_string_index = int(value)
        except ValueError:
            return value
        if shared_string_index < 0:
            return value
        try:
            return shared_strings[shared_string_index]
        except IndexError:
            return value
    return value.strip()


def _xml_local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _single_page_result(
    raw_text: str,
    *,
    page_count: int | None,
    notes: str | None,
    tables: list[ExtractedTable] | None = None,
) -> ExtractionResult:
    page = _make_page(
        raw_text,
        page_number=None,
        needs_ocr=False,
        source_span_start=0,
        notes=notes,
    )

    pages = _pages_with_raw_text(page)
    return _result_from_pages(
        pages,
        tables=tables,
        page_count=page_count,
        notes=notes,
    )


def _pages_with_raw_text(page: ExtractedPage) -> list[ExtractedPage]:
    if page.clean_text or page.raw_text.strip():
        return [page]
    return []
